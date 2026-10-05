"""The group order spreadsheet carries every order's full details."""
import io
import json
from datetime import datetime

import openpyxl

from models import db, Order, OrderItem
from tests.conftest import ADMIN_EMAIL, CUSTOMER_EMAIL


def _add_orders(app, seed):
    with app.app_context():
        shipped = Order(
            order_number='PM-SHIP-1', email='jo@example.com', first_name='Jo', last_name='Lane',
            phone='816-555-0111', collection_id=seed['collection_id'],
            fulfillment_method='shipping', shipping_recipient='Jo Lane',
            shipping_street='12 Elm St', shipping_city='Olathe', shipping_state='KS', shipping_zip='66061',
            subtotal=60, shipping_cost=11, tax=5.7, total=76.7, amount_paid=76.7,
            payment_method='stripe', payment_status='paid', status='paid',
            production_stage='ready_to_press', customer_notes='Gift, please wrap',
            admin_notes='Called about sizing', cost_of_goods=20, profit=56.7,
            paid_at=datetime(2026, 10, 1, 15, 0),
        )
        kid = Order(
            order_number='PM-KID-2', email='sam@example.com', first_name='Sam', last_name='Diaz',
            collection_id=seed['collection_id'], fulfillment_method='pickup',
            send_home_with_child=True, child_name='Mia Diaz', teacher_name='Mrs. Park', child_grade='2nd',
            subtotal=30, tax=2.85, total=32.85, payment_method='cash', payment_status='pending', status='new',
        )
        db.session.add_all([shipped, kid])
        db.session.flush()
        db.session.add_all([
            OrderItem(order_id=shipped.id, product_id=seed['tee_id'], product_name='Unisex Tee', style_number='3001',
                      size='M', color='Black', quantity=1, unit_price=30, subtotal=30,
                      placement='left_chest', design_file_name='Falcons Logo', print_width=4, print_height=3.5,
                      back_design_meta=json.dumps({'name': 'LANE', 'number': '12', 'font': 'Varsity',
                                                   'text_color': '#ffffff', 'outline': True, 'outline_color': '#000000'}),
                      notes='Extra long'),
            OrderItem(order_id=shipped.id, product_id=seed['tee_id'], product_name='Unisex Tee', style_number='3001',
                      size='L', color='Black', quantity=1, unit_price=30, subtotal=30, placement='center_chest'),
            OrderItem(order_id=kid.id, product_id=seed['youth_id'], product_name='Youth Tee', style_number='3001Y',
                      size='YM', color='White', quantity=1, unit_price=30, subtotal=30, placement='center_chest'),
        ])
        db.session.commit()


def _workbook(client):
    resp = client.get('/c/test-elementary/export.xlsx')
    assert resp.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(resp.data))


def _rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:]]


def test_all_orders_sheet_opens_first_with_every_detail(client, seed, app, login):
    _add_orders(app, seed)
    login(client, ADMIN_EMAIL)
    wb = _workbook(client)
    assert wb.sheetnames[0] == 'All Orders'
    assert wb.active.title == 'All Orders'
    rows = _rows(wb['All Orders'])
    assert len(rows) == 3  # one row per shirt

    first = next(r for r in rows if r['Order #'] == 'PM-SHIP-1' and r['Size'] == 'M')
    assert first['Customer'] == 'Jo Lane'
    assert first['Phone'] == '816-555-0111'
    assert first['Ship To'] == 'Jo Lane, 12 Elm St, Olathe, KS 66061'
    assert first['Front Design'] == 'Falcons Logo'
    assert first['Placement'] == 'Left chest'
    assert first['Print Size (in)'] == '4 x 3.5'
    assert first['Back Name'] == 'LANE'
    assert first['Back Number'] == '12'
    assert first['Back Font'] == 'Varsity'
    assert first['Item Notes'] == 'Extra long'
    assert first['Paid With'] == 'Card'
    assert first['Order Total'] == 76.7
    assert first['Production Stage'] == 'Ready To Press'
    assert first['Customer Notes'] == 'Gift, please wrap'
    assert first['Admin Notes'] == 'Called about sizing'

    # Order totals only on the order's first row so sums are not doubled.
    second = next(r for r in rows if r['Order #'] == 'PM-SHIP-1' and r['Size'] == 'L')
    assert second['Order Total'] in (None, '')
    assert second['Customer'] == 'Jo Lane'

    kid = next(r for r in rows if r['Order #'] == 'PM-KID-2')
    assert kid['Child'] == 'Mia Diaz'
    assert kid['Coach/Teacher'] == 'Mrs. Park'
    assert kid['Grade'] == '2nd'
    assert kid['Send Home'] == 'Yes'
    assert kid['Paid With'] == 'Cash'
    assert kid['Payment Status'] == 'Pending'


def test_orders_sheet_has_payment_and_shipping_details(client, seed, app, login):
    _add_orders(app, seed)
    login(client, ADMIN_EMAIL)
    rows = _rows(_workbook(client)['Orders'])
    shipped = next(r for r in rows if r['Order #'] == 'PM-SHIP-1')
    assert shipped['Paid With'] == 'Card'
    assert shipped['Amount Paid'] == 76.7
    assert shipped['Ship To'].startswith('Jo Lane, 12 Elm St')
    assert shipped['Profit'] == 56.7


def test_organizer_export_leaves_out_admin_only_columns(client, seed, app, login):
    from models import Collection
    _add_orders(app, seed)
    with app.app_context():
        db.session.get(Collection, seed['collection_id']).created_by_user_id = seed['customer_id']
        db.session.commit()
    login(client, CUSTOMER_EMAIL)
    wb = _workbook(client)
    assert 'Admin Notes' not in [c.value for c in wb['All Orders'][1]]
    assert 'Profit' not in [c.value for c in wb['Orders'][1]]
    assert 'Called about sizing' not in str(list(wb['All Orders'].values))


def test_spreadsheet_downloads_are_never_cached(client, seed, app, login):
    """Cloudflare caches .xlsx by default and served a stale order sheet."""
    login(client, ADMIN_EMAIL)
    resp = client.get('/c/test-elementary/export.xlsx')
    assert 'no-store' in resp.headers['Cache-Control']
    assert 'private' in resp.headers['Cache-Control']
    assert resp.headers['Cloudflare-CDN-Cache-Control'] == 'no-store'
    html = client.get('/admin/collections').get_data(as_text=True)
    assert '/c/test-elementary/export.xlsx?v=' in html
