"""Organizer-pays group orders: the link collects name + size, the organizer pays once."""
import io
import json

import openpyxl

from models import db, Collection, GroupRosterEntry, Order, OrderItem
from tests.conftest import ADMIN_EMAIL, CUSTOMER_EMAIL


def _in_db(app, fn):
    """Run fn() against the database in its own thread.

    The test clients here keep their last request context open, and opening
    an app context on this thread between requests makes Flask pop contexts
    in the wrong order. A worker thread has its own context stack.
    """
    import threading
    out = {}

    def run():
        with app.app_context():
            try:
                out['value'] = fn()
            except BaseException as exc:  # surface assertion errors to the test
                out['error'] = exc

    t = threading.Thread(target=run)
    t.start()
    t.join()
    if 'error' in out:
        raise out['error']
    return out.get('value')


def _make_roster_store(app, seed, *, expected=3, allow_cash=False, **extra):
    """Turn the seeded group store into an organizer-pays store: one tee, Black, one logo."""
    with app.app_context():
        c = db.session.get(Collection, seed['collection_id'])
        c.payment_mode = 'organizer_pays'
        c.expected_count = expected
        c.allowed_colors = json.dumps({f'p:{seed["tee_id"]}': ['Black']})
        c.allowed_design_ids = json.dumps([seed['free_design_id']])
        c.allowed_placements = json.dumps(['left_chest'])
        c.allow_cash_pickup = allow_cash
        for key, value in extra.items():
            setattr(c, key, value)
        db.session.commit()
    return seed['collection_slug']


def _submit(client, slug, first='Ava', last='Lopez', size='M', **extra):
    data = {'first_name': first, 'last_name': last, 'size': size}
    data.update(extra)
    return client.post(f'/c/{slug}/roster', data=data, follow_redirects=True)


# ── Organizer form ───────────────────────────────────────────────────────────

def test_create_form_asks_how_the_order_is_paid_for(customer_client):
    html = customer_client.get('/shop/group-orders/create').get_data(as_text=True)
    assert 'How will this order be paid for?' in html
    assert 'value="each_pays" checked' in html
    assert "I'm paying for the whole group" in html


def test_new_store_defaults_to_each_person_paying(customer_client, seed, app):
    customer_client.post(
        '/shop/group-orders/create',
        data={'name': 'Default Pay Store', 'products': [str(seed['tee_id'])], 'group_kind': 'other'},
        content_type='multipart/form-data', follow_redirects=True,
    )
    def _check1():
        c = Collection.query.filter_by(name='Default Pay Store').first()
        assert c is not None
        assert c.payment_mode == 'each_pays'
    _in_db(app, _check1)


def test_organizer_can_create_an_organizer_pays_store(customer_client, seed, app):
    resp = customer_client.post(
        '/shop/group-orders/create',
        data={
            'name': 'Riverview Falcons Spirit Wear',
            'payment_mode': 'organizer_pays',
            'expected_count': '20',
            'products': [str(seed['tee_id'])],
            'allowed_colors': [f'p:{seed["tee_id"]}||Black'],
            'allowed_designs': [str(seed['free_design_id'])],
            'allowed_placements': ['center_chest'],
            'group_kind': 'school',
        },
        content_type='multipart/form-data', follow_redirects=True,
    )
    html = resp.get_data(as_text=True)
    assert 'Sizes so far' in html
    assert '0 of 20' in html
    def _check2():
        c = Collection.query.filter_by(name='Riverview Falcons Spirit Wear').first()
        assert c.payment_mode == 'organizer_pays'
        assert c.expected_count == 20
    _in_db(app, _check2)


def test_organizer_pays_needs_exactly_one_shirt_and_color(customer_client, seed, app):
    html = customer_client.post(
        '/shop/group-orders/create',
        data={
            'name': 'Too Many Shirts',
            'payment_mode': 'organizer_pays',
            'products': [str(seed['tee_id']), str(seed['hoodie_id'])],
            'group_kind': 'other',
        },
        content_type='multipart/form-data', follow_redirects=True,
    ).get_data(as_text=True)
    assert 'choose exactly one shirt style' in html
    html = customer_client.post(
        '/shop/group-orders/create',
        data={
            'name': 'No Color Picked',
            'payment_mode': 'organizer_pays',
            'products': [str(seed['tee_id'])],
            'group_kind': 'other',
        },
        content_type='multipart/form-data', follow_redirects=True,
    ).get_data(as_text=True)
    assert 'choose exactly one color' in html
    def _check3():
        assert Collection.query.filter_by(name='Too Many Shirts').first() is None
        assert Collection.query.filter_by(name='No Color Picked').first() is None
    _in_db(app, _check3)


def test_store_with_shopper_orders_cannot_switch_to_organizer_pays(admin_client, seed, app):
    def _check4():
        db.session.add(Order(
            order_number='PM-TEST-1', email='a@example.com', first_name='A', last_name='B',
            subtotal=10, total=10, collection_id=seed['collection_id'],
        ))
        db.session.commit()
    _in_db(app, _check4)
    html = admin_client.post(
        f'/admin/collections/{seed["collection_id"]}/edit',
        data={
            'name': 'Test Elementary Spirit Wear',
            'payment_mode': 'organizer_pays',
            'products': [str(seed['tee_id'])],
            'allowed_colors': [f'p:{seed["tee_id"]}||Black'],
            'is_active': 'on',
        },
        follow_redirects=True,
    ).get_data(as_text=True)
    assert 'cannot switch' in html
    def _check5():
        assert db.session.get(Collection, seed['collection_id']).payment_mode != 'organizer_pays'
    _in_db(app, _check5)


# ── Parent form ──────────────────────────────────────────────────────────────

def test_parent_sees_one_short_form_with_no_shopping(guest, seed, app):
    slug = _make_roster_store(app, seed)
    resp = guest.get(f'/c/{slug}')
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert 'Test Elementary Spirit Wear' in html
    assert 'name="first_name"' in html and 'name="last_name"' in html
    assert '<select id="size" name="size"' in html
    assert 'Black' in html
    assert 'No account and no payment needed.' in html
    # No store grid, cart button or customize links for parents.
    assert 'Add to Cart' not in html
    assert '/shop/customize/' not in html
    with guest.session_transaction() as sess:
        assert 'collection_id' not in sess


def test_parent_submissions_show_up_for_the_organizer_in_order(guest, seed, app, login):
    admin_client = guest
    slug = _make_roster_store(app, seed)
    html = _submit(guest, slug, 'Ava', 'Lopez', 'M').get_data(as_text=True)
    assert 'Thank you, Ava Lopez!' in html
    assert 'size M' in html
    _submit(guest, slug, 'Ben', "O'Brien", 'L')
    _submit(guest, slug, 'Cara', 'Ng', 'M')

    def _check6():
        rows = GroupRosterEntry.query.filter_by(collection_id=seed['collection_id']).order_by(GroupRosterEntry.id).all()
        assert [(r.first_name, r.last_name, r.size) for r in rows] == [
            ('Ava', 'Lopez', 'M'), ('Ben', "O'Brien", 'L'), ('Cara', 'Ng', 'M'),
        ]
    _in_db(app, _check6)

    login(admin_client, ADMIN_EMAIL)
    share = admin_client.get(f'/c/{slug}/share').get_data(as_text=True)
    assert '<strong>3 of 3</strong> responded' in share
    assert share.index('Ava Lopez') < share.index('Ben O') < share.index('Cara Ng')
    assert '<span>M</span> 2' in share and '<span>L</span> 1' in share
    assert 'Pay for the group (3 shirts)' in share


def test_parent_form_rejects_missing_name_or_unknown_size(guest, seed, app):
    slug = _make_roster_store(app, seed)
    resp = _submit(guest, slug, first='', size='M')
    assert resp.status_code == 400
    assert 'Please enter a first and last name.' in resp.get_data(as_text=True)
    resp = _submit(guest, slug, size='XXXXL')
    assert resp.status_code == 400
    assert 'Please choose a size from the list.' in resp.get_data(as_text=True)
    def _check7():
        assert GroupRosterEntry.query.count() == 0
    _in_db(app, _check7)


def test_honeypot_submission_is_dropped(guest, seed, app):
    slug = _make_roster_store(app, seed)
    _submit(guest, slug, website='spam.example')
    def _check8():
        assert GroupRosterEntry.query.count() == 0
    _in_db(app, _check8)


def test_parent_cannot_check_out_an_organizer_pays_store(guest, seed, app):
    slug = _make_roster_store(app, seed)
    guest.get(f'/c/{slug}')
    with guest.session_transaction() as sess:
        sess['collection_id'] = seed['collection_id']
    resp = guest.post('/cart/add', data={
        'product_id': seed['tee_id'], 'size': 'M', 'color': 'Black', 'quantity': 1,
    })
    assert resp.status_code == 400
    assert 'paid for by the organizer' in resp.get_json()['error']


def test_closed_form_after_deadline(guest, seed, app):
    from datetime import datetime, timedelta
    slug = _make_roster_store(app, seed, order_deadline=datetime.utcnow() - timedelta(days=2))
    html = guest.get(f'/c/{slug}').get_data(as_text=True)
    assert 'The deadline to send in sizes has passed.' in html
    assert 'name="first_name"' not in html
    _submit(guest, slug)
    def _check9():
        assert GroupRosterEntry.query.count() == 0
    _in_db(app, _check9)


# ── Organizer tools ──────────────────────────────────────────────────────────

def test_roster_spreadsheet_download(guest, seed, app, login):
    admin_client = guest
    slug = _make_roster_store(app, seed)
    _submit(guest, slug, 'Ava', 'Lopez', 'M')
    _submit(guest, slug, 'Ben', 'Ray', 'S')
    login(admin_client, ADMIN_EMAIL)
    resp = admin_client.get(f'/c/{slug}/roster.xlsx')
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    rows = list(wb['Roster'].iter_rows(values_only=True))
    assert rows[0][:4] == ('#', 'First name', 'Last name', 'Size')
    assert rows[1][1:4] == ('Ava', 'Lopez', 'M')
    assert rows[2][1:4] == ('Ben', 'Ray', 'S')
    totals = dict(list(wb['Size totals'].iter_rows(values_only=True))[1:])
    assert totals == {'S': 1, 'M': 1, 'Total': 2}


def test_strangers_cannot_see_or_download_the_roster(guest, seed, app, login):
    customer_client = guest
    slug = _make_roster_store(app, seed)
    _submit(guest, slug)
    login(customer_client, CUSTOMER_EMAIL)
    assert customer_client.get(f'/c/{slug}/roster.xlsx').status_code == 404
    share = customer_client.get(f'/c/{slug}/share').get_data(as_text=True)
    assert 'Ava Lopez' not in share


def test_organizer_can_remove_a_duplicate(guest, seed, app, login):
    admin_client = guest
    slug = _make_roster_store(app, seed)
    _submit(guest, slug)
    _submit(guest, slug)
    def _check10():
        return GroupRosterEntry.query.order_by(GroupRosterEntry.id.desc()).first().id
    dup = _in_db(app, _check10)
    login(admin_client, ADMIN_EMAIL)
    admin_client.post(f'/c/{slug}/roster/{dup}/remove', follow_redirects=True)
    def _check11():
        assert GroupRosterEntry.query.count() == 1
    _in_db(app, _check11)


def test_organizer_pays_once_for_the_whole_roster(guest, seed, app, login):
    admin_client = guest
    slug = _make_roster_store(app, seed, allow_cash=True)
    _submit(guest, slug, 'Ava', 'Lopez', 'M')
    _submit(guest, slug, 'Ben', 'Ray', 'L')
    _submit(guest, slug, 'Cara', 'Ng', 'M')

    login(admin_client, ADMIN_EMAIL)
    resp = admin_client.post(f'/c/{slug}/roster/pay')
    assert resp.status_code == 302
    assert '/checkout' in resp.headers['Location']
    assert admin_client.get('/checkout/').status_code == 200

    body = admin_client.post('/checkout/complete', json={
        'payment_method': 'cash', 'shipping_method': 'pickup',
        'email': 'admin-test@example.com', 'first_name': 'Site', 'last_name': 'Owner',
        'phone': '816-555-0100', 'checkout_token': 'tok-roster-1',
    }).get_json()
    assert body['success'], body

    def _check12():
        order = Order.query.filter_by(order_number=body['order_number']).first()
        assert order.collection_id == seed['collection_id']
        by_size = {i.size: i.quantity for i in OrderItem.query.filter_by(order_id=order.id)}
        assert by_size == {'M': 2, 'L': 1}
        assert all(i.color == 'Black' for i in order.items)
    _in_db(app, _check12)

    # Link closes once the order is placed (for parents too).
    admin_client.get('/auth/logout')
    html = guest.get(f'/c/{slug}').get_data(as_text=True)
    assert 'This order is closed.' in html
    _submit(guest, slug, 'Late', 'Parent', 'S')
    def _check13():
        assert GroupRosterEntry.query.count() == 3
    _in_db(app, _check13)
    login(admin_client, ADMIN_EMAIL)
    share = admin_client.get(f'/c/{slug}/share').get_data(as_text=True)
    assert 'has been placed' in share


def test_organizer_can_pay_after_the_size_deadline(guest, seed, app, login):
    admin_client = guest
    from datetime import datetime, timedelta
    slug = _make_roster_store(app, seed)
    _submit(guest, slug)
    def _check14():
        c = db.session.get(Collection, seed['collection_id'])
        c.order_deadline = datetime.utcnow() - timedelta(days=1)
        db.session.commit()
    _in_db(app, _check14)
    login(admin_client, ADMIN_EMAIL)
    resp = admin_client.post(f'/c/{slug}/roster/pay')
    assert '/checkout' in resp.headers['Location']
    assert admin_client.get('/checkout/').status_code == 200


# ── At-a-glance labels ───────────────────────────────────────────────────────

def test_admin_and_account_lists_show_the_payment_mode(admin_client, seed, app):
    html = admin_client.get('/admin/collections').get_data(as_text=True)
    assert 'Each person pays' in html
    _make_roster_store(app, seed)
    html = admin_client.get('/admin/collections').get_data(as_text=True)
    assert 'Organizer pays' in html
    mine = admin_client.get('/account/group-orders').get_data(as_text=True)
    assert 'You pay for the group' in mine
    assert '0 of 3 responded' in mine


def test_each_pays_store_is_unchanged(guest, seed):
    html = guest.get(f'/c/{seed["collection_slug"]}').get_data(as_text=True)
    assert 'name="first_name"' not in html
    assert 'Send my size' not in html
    with guest.session_transaction() as sess:
        assert sess.get('collection_id') == seed['collection_id']
    resp = guest.post('/cart/add', data={
        'product_id': seed['tee_id'], 'size': 'M', 'color': 'Black', 'quantity': 1,
        'placement': 'center_chest', 'design_id': seed['free_design_id'],
    })
    assert resp.status_code == 200
