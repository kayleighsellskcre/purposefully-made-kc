"""Organizer-pays group orders: the link collects name + size, the organizer pays once.

Every group order has a payment mode:

* ``each_pays`` (default, and what every older group order uses): shoppers
  open the link, customize, and check out on their own.
* ``organizer_pays``: the organizer picks one shirt, one color and the logo
  up front. The shared link is a one-screen form (first name, last name,
  size). Each submission is one roster line. When the organizer is ready
  they pay once for the whole roster through the normal checkout.

Nothing here changes the each-pays flow.
"""
from __future__ import annotations

import io
from collections import OrderedDict

EACH_PAYS = 'each_pays'
ORGANIZER_PAYS = 'organizer_pays'
PAYMENT_MODES = (EACH_PAYS, ORGANIZER_PAYS)

NAME_MAX = 80
EXPECTED_MAX = 5000
FRONT_PLACEMENTS = ('center_chest', 'left_chest', 'right_chest')


def is_organizer_pays(collection) -> bool:
    return bool(collection) and getattr(collection, 'payment_mode', None) == ORGANIZER_PAYS


def payment_mode_label(collection) -> str:
    """Short label for admin and account lists."""
    return 'Organizer pays' if is_organizer_pays(collection) else 'Each person pays'


def apply_payment_mode_from_form(collection):
    """Save the payment choice + expected headcount from the current request.

    Returns (ok, error_message). A missing field keeps the current mode, so
    forms that do not show the choice can never flip a store by accident.
    """
    from flask import request

    raw = (request.form.get('payment_mode') or '').strip()
    if raw:
        if raw not in PAYMENT_MODES:
            return False, 'Please choose how this group order will be paid for.'
        previous = getattr(collection, 'payment_mode', None) or EACH_PAYS
        if (
            raw == ORGANIZER_PAYS
            and previous != ORGANIZER_PAYS
            and getattr(collection, 'id', None)
            and collection.orders.count() > 0
        ):
            return False, ('Shoppers have already placed orders on this group order, so it '
                           'cannot switch to you paying for the whole group. Start a new '
                           'group order for that instead.')
        collection.payment_mode = raw
    elif not getattr(collection, 'payment_mode', None):
        collection.payment_mode = EACH_PAYS

    if 'expected_count' in request.form:
        count_raw = (request.form.get('expected_count') or '').strip()
        if not count_raw:
            collection.expected_count = None
        else:
            try:
                count = int(count_raw)
            except ValueError:
                return False, 'Please enter how many people you expect as a whole number.'
            if count < 1 or count > EXPECTED_MAX:
                return False, 'Please enter how many people you expect (1 or more).'
            collection.expected_count = count
    return True, None


# ── The one item everyone gets ──────────────────────────────────────────────

def _fan_products(collection):
    from utils.group_orders import team_store_config, visible_store_products

    config = team_store_config(collection)
    products = visible_store_products(collection)
    if config['configured']:
        fan_ids = set(config['fan_product_ids'])
        products = [p for p in products if p.id in fan_ids]
    return products, config


def _pick_design_id(collection):
    from utils.group_orders import allowed_design_ids, showcase_design_ids

    allowed = allowed_design_ids(collection)
    for did in showcase_design_ids(collection):
        if did in allowed:
            return did
    return allowed[0] if allowed else None


def roster_item_error(collection):
    """Why this store cannot run as organizer-pays yet, or None.

    Checked when the organizer saves, so the share link always has exactly
    one shirt in one color to show parents.
    """
    if not is_organizer_pays(collection):
        return None
    from utils.group_orders import allowed_colors_for_product

    products, config = _fan_products(collection)
    if config['uniform']['enabled']:
        return ('When you are paying for the whole group, turn off the player uniform '
                'and choose one shirt instead.')
    if len(products) != 1:
        return ('When you are paying for the whole group, choose exactly one shirt style. '
                f'You have {len(products)} selected.')
    allowed = allowed_colors_for_product(products[0], collection)
    if not allowed or len(allowed) != 1:
        return ('When you are paying for the whole group, choose exactly one color '
                'for the shirt so everyone gets the same thing.')
    return None


def extra_logo_count(collection) -> int:
    from utils.group_orders import allowed_design_ids
    return max(0, len(allowed_design_ids(collection)) - 1)


def roster_item(collection, app=None):
    """Everything the parent form and the bulk checkout need, or None.

    Returns a dict: product, color, design (dict or None), design_id,
    placement, back_design (dict or None), image_url, sizes, size_chart,
    spec_sheet_url.
    """
    if roster_item_error(collection):
        return None
    from flask import current_app
    from utils.group_orders import allowed_colors_for_product, load_design_dict
    from utils.json_fields import parse_json_list
    from utils.mockups import get_carousel_colors_for_product, get_first_shop_image_url
    from utils.sizes import sort_sizes

    app = app or current_app
    products, _config = _fan_products(collection)
    product = products[0]
    color = next(iter(allowed_colors_for_product(product, collection)))

    design_id = _pick_design_id(collection)
    design = load_design_dict(design_id) if design_id else None
    if not design:
        design_id = None

    placement = None
    if design:
        placements = [
            (p or '').strip().lower()
            for p in parse_json_list(getattr(collection, 'allowed_placements', None) or '')
        ]
        front = [p for p in placements if p in FRONT_PLACEMENTS]
        placement = front[0] if front else 'center_chest'

    back = None
    if (
        getattr(collection, 'allow_back_design', True)
        and (getattr(collection, 'back_design_type', None) or 'both') == 'image'
        and getattr(collection, 'back_design_id', None)
    ):
        back = load_design_dict(collection.back_design_id)

    variants = get_carousel_colors_for_product(product, app, allowed_colors={color})
    image_url = None
    if variants:
        image_url = variants[0].get('front_image_url')
    if not image_url:
        image_url = get_first_shop_image_url(product, app, carousel=variants)

    size_chart = None
    raw_chart = getattr(product, 'size_chart', None)
    if raw_chart:
        import json
        try:
            parsed = json.loads(raw_chart)
            size_chart = parsed if isinstance(parsed, dict) and parsed else None
        except (TypeError, ValueError):
            size_chart = None

    return {
        'product': product,
        'color': color,
        'design': design,
        'design_id': design_id,
        'placement': placement,
        'back_design': back,
        'image_url': image_url,
        'sizes': sort_sizes(parse_json_list(product.available_sizes)),
        'size_chart': size_chart,
        'spec_sheet_url': getattr(product, 'spec_sheet_url', None),
    }


# ── Roster ──────────────────────────────────────────────────────────────────

def roster_closed(collection) -> bool:
    """The organizer has placed the group's order, so the link stops taking sizes."""
    return collection.orders.count() > 0


def roster_entries(collection):
    from models import GroupRosterEntry
    return (
        GroupRosterEntry.query
        .filter_by(collection_id=collection.id)
        .order_by(GroupRosterEntry.created_at, GroupRosterEntry.id)
        .all()
    )


def size_counts(entries, sizes=None):
    """OrderedDict size -> count, in shirt-size order."""
    from utils.sizes import sort_sizes

    counts = {}
    for e in entries:
        counts[e.size] = counts.get(e.size, 0) + 1
    order = sort_sizes(list(counts.keys()))
    return OrderedDict((s, counts[s]) for s in order)


def roster_summary(collection):
    entries = roster_entries(collection)
    return {
        'entries': entries,
        'count': len(entries),
        'expected': getattr(collection, 'expected_count', None),
        'size_counts': size_counts(entries),
        'closed': roster_closed(collection),
    }


def validate_submission(form, item):
    """Clean (first, last, size) from a parent's form, or an error message."""
    first = ' '.join((form.get('first_name') or '').split())
    last = ' '.join((form.get('last_name') or '').split())
    size = (form.get('size') or '').strip()
    if not first or not last:
        return None, 'Please enter a first and last name.'
    if len(first) > NAME_MAX or len(last) > NAME_MAX:
        return None, 'Please keep each name under 80 characters.'
    if size not in (item.get('sizes') or []):
        return None, 'Please choose a size from the list.'
    return (first, last, size), None


# ── Organizer pays once ─────────────────────────────────────────────────────

def build_roster_cart(collection, item, entries):
    """Cart lines for the whole roster: one line per size, quantity = headcount."""
    from models import Design
    from utils.pricing import calculate_unit_price
    from utils.print_sizes import build_item_production

    product = item['product']
    design = item.get('design')
    design_row = Design.query.get(item['design_id']) if item.get('design_id') else None
    back = item.get('back_design')
    design_fee = float(getattr(design_row, 'design_fee', 0) or 0) if design_row else 0.0
    aspect_w = float(design_row.width) if design_row and design_row.width else None
    aspect_h = float(design_row.height) if design_row and design_row.height else None
    design_name = (design or {}).get('title')

    lines = []
    for size, qty in size_counts(entries).items():
        unit_price = calculate_unit_price(
            product,
            size=size,
            placement=item.get('placement'),
            has_back_design=bool(back),
            is_blank=not (design or back),
            design_fee=design_fee,
        )
        production = build_item_production(
            product=product,
            size=size,
            color=item['color'],
            placement=item.get('placement'),
            quantity=qty,
            design_name=design_name,
            design_id=item.get('design_id'),
            aspect_w=aspect_w,
            aspect_h=aspect_h,
            has_front=bool(design),
            customer_name=collection.name,
        )
        front = (production or {}).get('front') or {}
        lines.append({
            'product_id': product.id,
            'size': size,
            'color': item['color'],
            'quantity': qty,
            'unit_price': unit_price,
            'is_blank': not (design or back),
            'design_id': item.get('design_id'),
            'design_url': (design or {}).get('url'),
            'placement': item.get('placement'),
            'back_design_url': (back or {}).get('url'),
            'back_design_kind': 'image' if back else None,
            'back_design_meta': None,
            'proof_front_url': None,
            'proof_back_url': None,
            'print_width': front.get('width'),
            'print_height': front.get('height'),
            'transfer_production': production,
            'position_x': None,
            'position_y': None,
            'rotation': 0,
            'proof_image': None,
            'collection_id': collection.id,
            'catalog_section': 'fan',
            'uniform_kit': None,
            'roster_order': True,
        })
    return lines


def roster_xlsx(collection) -> bytes:
    """Spreadsheet of the roster: names + sizes in the order submitted, plus totals."""
    import openpyxl
    from openpyxl.styles import Font
    from utils.local_time import format_central

    summary = roster_summary(collection)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Roster'
    ws.append(['#', 'First name', 'Last name', 'Size', 'Submitted'])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for n, e in enumerate(summary['entries'], start=1):
        ws.append([
            n, e.first_name, e.last_name, e.size,
            format_central(e.created_at, '%b %d, %Y %I:%M %p') if e.created_at else '',
        ])
    for col, width in zip('ABCDE', (5, 18, 18, 8, 22)):
        ws.column_dimensions[col].width = width

    totals = wb.create_sheet('Size totals')
    totals.append(['Size', 'Count'])
    for cell in totals[1]:
        cell.font = Font(bold=True)
    for size, count in summary['size_counts'].items():
        totals.append([size, count])
    totals.append(['Total', summary['count']])
    totals.cell(row=totals.max_row, column=1).font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
